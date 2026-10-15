# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.
#
# Copyright (c) 2026 Zarrakun

import argparse
import datetime
import os
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from .probe import ENGINE_ROOT, collect_env, run_checks
RUNS_HEADERS = ["RunID", "Date", "Tester", "Machine", "GPU", "Arch", "Driver", "OS",
                "CPU", "RAM_GB", "Engine_Commit", "Branch", "Python", "PyQt", "moderngl",
                "Scene", "FPS_avg", "FPS_1pct", "FPS_min", "Result", "Issues", "Log", "Notes"]
CHECKS_HEADERS = ["RunID", "Area", "Check", "Expected", "Actual", "Status"]
GPUS_HEADERS = ["GPU", "Vendor", "Arch", "VRAM", "Machine", "Driver_Pinned", "OS", "Owner", "Status", "Notes"]
SCENES_HEADERS = ["Scene", "Path", "Purpose", "Min_FPS_Target"]
FLEET = [
    ["RTX 3090", "NVIDIA", "Ampere", "24GB", "", "", "Windows", "", "Owned", "Reference rig, everything must pass here"],
    ["Radeon RX 580", "AMD", "GCN Polaris", "8GB", "", "", "Windows", "", "Owned", "Strict GLSL compiler, catches sampler/index bugs"],
    ["GTX 1050 Ti", "NVIDIA", "Pascal", "4GB", "", "", "Windows", "", "Planned", "Mass-market floor, perf and VRAM pressure"],
    ["RX 5500 XT", "AMD", "RDNA1", "4/8GB", "", "", "Windows", "", "Planned", "Second AMD compiler backend"],
    ["UHD 630", "Intel", "Gen9.5", "shared", "", "", "Windows", "", "Planned", "Weakest GL compliance, scrap office PC is fine"],
    ["Arc A380", "Intel", "Xe-HPG", "6GB", "", "", "Windows", "", "Planned", "Young driver stack, new bug classes"],
]
SCENES_SEED = [
    ["grid_empty", "", "Grid rendering, LOD fade, no shimmer", "60"],
    ["sun_floor", "", "Directional shadows, cascades, atlas sampling", "60"],
    ["lights_torture", "", "Point and spot shadows, sampler budget under load", "30"],
]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
def style_sheet(ws, headers, widths):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr = None
def add_validation(ws, col_letter, formula, start_row=2, end_row=5000):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Pick from the list"
    dv.errorTitle = "Invalid value"
    ws.add_data_validation(dv)
    dv.add(col_letter + str(start_row) + ":" + col_letter + str(end_row))
def ensure_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Runs"
    style_sheet(ws, RUNS_HEADERS, [8, 16, 14, 14, 28, 14, 22, 16, 28, 8, 14, 8, 8, 8, 10, 18, 9, 9, 9, 10, 20, 20, 30])
    add_validation(ws, "T", '"PASS,FAIL,BLOCKED"')
    ws2 = wb.create_sheet("Checks")
    style_sheet(ws2, CHECKS_HEADERS, [8, 16, 44, 30, 60, 10])
    add_validation(ws2, "F", '"PASS,FAIL,SKIP"')
    ws3 = wb.create_sheet("GPUs")
    style_sheet(ws3, GPUS_HEADERS, [22, 10, 14, 8, 14, 18, 12, 12, 10, 40])
    for row in FLEET:
        ws3.append(row)
    for r in range(2, 2 + len(FLEET)):
        for c in range(1, len(GPUS_HEADERS) + 1):
            ws3.cell(row=r, column=c).border = BORDER
    add_validation(ws3, "I", '"Owned,Planned,Retired"')
    ws4 = wb.create_sheet("Scenes")
    style_sheet(ws4, SCENES_HEADERS, [24, 40, 52, 14])
    for row in SCENES_SEED:
        ws4.append(row)
    for r in range(2, 2 + len(SCENES_SEED)):
        for c in range(1, len(SCENES_HEADERS) + 1):
            ws4.cell(row=r, column=c).border = BORDER
    wb.save(path)
    return path
def append_run(path, env, rows):
    wb = load_workbook(path)
    ws = wb["Runs"]
    run_id = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, int) and v >= run_id:
            run_id = v + 1
    record = [run_id, env["date"], env["tester"], env["machine"], env["gpu"], env["arch"],
              env["driver"], env["os"], env["cpu"], env["ram"], env["commit"], env["branch"],
              env["python"], env["pyqt"], env["moderngl"], env["scene"], "", "", "",
              "FAIL" if any(s == "FAIL" for _, _, _, _, s in rows) else "PASS",
              "", "", ""]
    ws.append(record)
    for c in range(1, len(record) + 1):
        ws.cell(row=ws.max_row, column=c).border = BORDER
    for c in [7, 11, 12, 13, 14, 15]:
        ws.cell(row=ws.max_row, column=c).number_format = "@"
    ws2 = wb["Checks"]
    for area, name, exp, act, status in rows:
        ws2.append([run_id, area, name, exp, act, status])
        for c in range(1, 7):
            ws2.cell(row=ws2.max_row, column=c).border = BORDER
    wb.save(path)
    return run_id
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tester", default="")
    ap.add_argument("--machine", default="")
    ap.add_argument("--arch", default="")
    ap.add_argument("--scene", default="")
    ap.add_argument("--journal", default=os.path.join(ENGINE_ROOT, "tests", "TestJournal.xlsx"))
    ap.add_argument("--no-save", action="store_true")
    args = ap.parse_args()
    env = collect_env(args)
    if args.machine:
        env["machine"] = args.machine
    print("GPU: " + (env["gpu"] or "?") + " | driver: " + (env["driver"] or "?") + " | commit: " + (env["commit"] or "?"))
    if env["gl_error"]:
        print("GL unavailable: " + env["gl_error"])
    rows = run_checks(env)
    for area, name, exp, act, status in rows:
        print("[" + status + "] " + area + " / " + name + " :: " + act)
    if not args.no_save:
        if not os.path.exists(args.journal):
            ensure_workbook(args.journal)
            print("journal created: " + args.journal)
        try:
            run_id = append_run(args.journal, env, rows)
            print("run " + chr(35) + str(run_id) + " saved to " + args.journal)
        except PermissionError:
            base, ext = os.path.splitext(args.journal)
            alt = base + "_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S") + ext
            import shutil
            shutil.copyfile(args.journal, alt)
            run_id = append_run(alt, env, rows)
            print("journal is open in Excel, run " + chr(35) + str(run_id) + " saved to " + alt)
    bad = sum(1 for _, _, _, _, s in rows if s == "FAIL")
    return 1 if bad else 0
if __name__ == "__main__":
    sys.exit(main())
